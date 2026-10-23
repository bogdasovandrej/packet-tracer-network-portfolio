from pathlib import Path
from subprocess import run
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
IMG = ROOT / 'figures'
IMG.mkdir(exist_ok=True)


def font(size=23, mono=False):
    path = Path('C:/Windows/Fonts/consola.ttf' if mono else 'C:/Windows/Fonts/arial.ttf')
    if path.exists():
        return ImageFont.truetype(str(path), size)
    fallback = 'DejaVuSansMono.ttf' if mono else 'DejaVuSans.ttf'
    try:
        return ImageFont.truetype(fallback, size)
    except OSError:
        return ImageFont.load_default()


def terminal(name, command):
    result = run([sys.executable, *command], cwd=ROOT, capture_output=True, text=True, check=True)
    lines = [f'PS quantum-telecom-labs> python {" ".join(command)}', *result.stdout.strip().splitlines()]
    lines = [line.replace('Alice received key ID:', 'Alice received key ID:') for line in lines]
    if name == 'api':
        lines = [line[:23] + ' <UUID> HTTP 200' if line.startswith('Alice received key ID:') else line for line in lines]
    width = max(1100, max(len(line) for line in lines) * 14 + 70)
    im = Image.new('RGB', (width, len(lines)*35+50), '#14202e')
    d = ImageDraw.Draw(im)
    f = font(22, True)
    for i, line in enumerate(lines):
        d.text((25, 22+i*35), line, fill='#edf4fa' if i else '#7fe0b2', font=f)
    path = IMG / f'{name}-terminal.png'
    im.save(path)
    return path


def chart(name, title, points, ylabel, color):
    w, h = 1100, 590
    im = Image.new('RGB', (w,h), 'white')
    d = ImageDraw.Draw(im)
    d.text((62,25), title, fill='#132b43', font=font(29))
    x0, y0, x1, y1 = 95, 95, 1010, 490
    d.line((x0,y0,x0,y1,x1,y1), fill='#52677b', width=3)
    vmax = max(y for _,y in points)*1.12
    for i in range(5):
        yy = y1 - (y1-y0)*i/4
        d.line((x0,yy,x1,yy), fill='#dce4e9', width=1)
        d.text((5,yy-12), f'{vmax*i/4:.1f}', fill='#445b6e', font=font(20))
    xs = [x for x,_ in points]
    coords = []
    for x,y in points:
        px = x0 + (x-min(xs))/(max(xs)-min(xs))*(x1-x0)
        py = y1 - y/vmax*(y1-y0)
        coords.append((px,py))
        d.ellipse((px-6,py-6,px+6,py+6), fill=color)
        d.text((px-12,y1+12), str(x), fill='#445b6e', font=font(19))
    d.line(coords, fill=color, width=4)
    d.text((480,535),'Длина волокна, км',fill='#445b6e',font=font(21))
    d.text((95,60),ylabel,fill='#445b6e',font=font(20))
    path = IMG / f'{name}.png'
    im.save(path)
    return path


def setup_doc(title, subtitle):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.68)
    sec.left_margin = sec.right_margin = Inches(.78)
    styles = doc.styles
    styles['Normal'].font.name = 'Aptos'
    styles['Normal'].font.size = Pt(10.5)
    styles['Normal'].paragraph_format.space_after = Pt(6)
    for name, size in [('Title', 24), ('Heading 1', 15), ('Heading 2', 11.5)]:
        styles[name].font.name = 'Aptos Display' if name == 'Title' else 'Aptos'
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor(20,43,67)
        styles[name].font.bold = name != 'Title'
    styles['Title'].font.color.rgb = RGBColor(0,0,0)
    p=doc.add_paragraph(style='Title'); p.add_run(title)
    border=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom')
    bottom.set(qn('w:val'),'nil'); border.append(bottom); p._p.get_or_add_pPr().append(border)
    p=doc.add_paragraph(subtitle); p.style='Subtitle'
    return doc


def h(doc, text, level=1): doc.add_heading(text, level)
def p(doc, text): doc.add_paragraph(text)
def b(doc, text): doc.add_paragraph(text, style='List Bullet')
def n(doc, text): doc.add_paragraph(text, style='List Number')


def code(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(.25)
    para.paragraph_format.space_after = Pt(8)
    para.paragraph_format.keep_together = True
    for line in text.strip().splitlines():
        r=para.add_run(line+'\n'); r.font.name='Consolas'; r.font.size=Pt(8.5)


def picture(doc, path, caption):
    para=doc.add_paragraph()
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(path), width=Inches(6.75))
    cp=doc.add_paragraph(caption)
    cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cp.runs[0].italic=True
    cp.runs[0].font.size=Pt(8.5)


def table(doc, headers, rows):
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Light Shading Accent 1'
    for c,s in zip(t.rows[0].cells,headers): c.text=s
    for row in rows:
        for c,s in zip(t.add_row().cells,row): c.text=str(s)


def sources(doc, items):
    h(doc,'Источники и честные ограничения')
    for item in items: b(doc,item)


opt_terminal=terminal('optical',['-m','lab01.link_budget'])
bb_terminal=terminal('bb84',['-m','lab02.bb84','--pulses','100000','--intercept'])
api_terminal=terminal('api',['-m','lab03.demo'])
opt_chart=chart('optical-clicks','Регистрации падают с расстоянием',[(0,15058),(20,6022.7),(40,2402.6),(60,957.8),(80,382),(100,152.7)],'Регистрации в секунду','#137b9e')
qber_chart=chart('optical-qber','Доля ошибочных регистраций растёт',[(0,2),(20,2.01),(40,2.02),(60,2.05),(80,2.13),(100,2.31)],'QBER, %','#c36a22')


def lab1():
    d=setup_doc('Оптическая линия и бюджет потерь','Лабораторная 1 · инженерная оценка квантового канала · Python')
    p(d,'Цель: научиться считать потери оптической линии, отличать реальные потери от инженерного запаса и связывать ослабление сигнала с регистрациями и QBER. В конце ты сможешь объяснить, почему канал, работающий на короткой дистанции, может стать ненадёжным после удлинения волокна.')
    h(d,'Что потребуется')
    b(d,'Windows PowerShell и Python 3.10 или новее. Пакеты устанавливать не нужно.')
    b(d,'Файл lab01/link_budget.py из каталога quantum-telecom-labs; изменения делай в собственной ветке или копии.')
    h(d,'Шаг 1 Запусти исходную модель')
    n(d,'Открой PowerShell в каталоге quantum-telecom-labs. Проверь python --version. Команда показывает выбранный интерпретатор; если Python не найден, установи его и вновь открой терминал.')
    code(d,'python --version\npython -m lab01.link_budget')
    p(d,'Ключ -m просит Python найти модуль lab01.link_budget внутри текущего каталога и выполнить блок if __name__ == "__main__". Не запускай файл из другой папки: тогда Python может не найти пакет.')
    picture(d,opt_terminal,'Снимок реального вывода программы для параметров по умолчанию')
    h(d,'Шаг 2 Разбери физический бюджет')
    p(d,'Длина 40 км при затухании 0,2 дБ/км даёт 8 дБ. Два разъёма по 0,5 дБ дают 1 дБ, две сварки по 0,1 дБ — ещё 0,2 дБ. Итого фактическая модельная потеря 9,2 дБ. Запас 3 дБ добавляют при планировании: 12,2 дБ. Он не является ещё одним затуханием в уже описанной физической линии.')
    code(d,'L_физ = расстояние × затухание + N_разъёмов × L_разъёма + N_сварок × L_сварки\nL_план = L_физ + инженерный_запас\nT = 10 ** (-L_физ / 10)')
    p(d,'T — доля оптической мощности после линии. Децибелы складывают, а коэффициенты передачи перемножают. Например, 10 дБ ослабления оставляют примерно одну десятую мощности.')
    h(d,'Шаг 3 Свяжи потери с регистрацией')
    p(d,'Учебная модель считает вероятность сигнального клика как 1 − exp(−μTη), где μ — среднее число фотонов на импульс, T — передача линии, η — эффективность детектора. Затем добавляет вероятность тёмного клика, то есть регистрации без сигнального фотона. Это приближение с пуассоновским источником; оно не описывает все детали реального QKD-приёмника.')
    code(d,'P_сигн = 1 − exp(−μ × T × η)\nP_клик = P_сигн + (1 − P_сигн) × P_тёмн\nR_клик = частота_импульсов × P_клик')
    p(d,'При 40 км пример даёт около 2403 регистраций в секунду, а при 100 км — около 153. Это сырые регистрации, не число секретных битов: часть отбрасывают при просеивании, коррекции и усилении секретности.')
    picture(d,opt_chart,'Ось Y показывает модельную частоту сырых регистраций, а не скорость секретного ключа')
    h(d,'Шаг 4 Объясни QBER')
    p(d,'Для сигнальных кликов берём ошибку оптического тракта 2%. Тёмный клик случайно даёт неверный бит примерно в половине случаев. Когда сигнал ослабевает, доля тёмных кликов растёт, поэтому растёт и QBER — отношение ошибочных битов к сравнимым битам. В этой модели QBER не включает последствия всех атак и не доказывает безопасность.')
    picture(d,qber_chart,'Даже при одинаковой оптике слабый сигнал увеличивает относительный вклад шумов')
    h(d,'Шаг 5 Измени один параметр и проверь гипотезу')
    p(d,'1. В PowerShell запусти следующую команду: она создаёт несколько объектов Link, меняя только длину, и печатает скорость кликов и QBER.')
    code(d,'python -c "from lab01.link_budget import Link; print([(x, round(Link(distance_km=x).raw_clicks_per_second), round(100*Link(distance_km=x).qber_estimate, 2)) for x in (10, 50, 90)])"')
    p(d,'2. До запуска запиши прогноз: при переходе 10 → 90 км число кликов уменьшится, QBER увеличится. Проверь вывод и объясни его через T и долю тёмных кликов.')
    p(d,'3. Сравни два изменения по отдельности на дистанции 40 км. Первая строка увеличивает эффективность детектора, вторая — вероятность тёмного клика. Сравни с исходной строкой 40 км из таблицы: не меняй два параметра одновременно, иначе причину различия трудно установить.')
    code(d,'python -c "from lab01.link_budget import Link; x=Link(distance_km=40,detector_efficiency=0.4); print(round(x.raw_clicks_per_second,1),round(100*x.qber_estimate,2))"\npython -c "from lab01.link_budget import Link; x=Link(distance_km=40,dark_click_probability=1e-5); print(round(x.raw_clicks_per_second,1),round(100*x.qber_estimate,2))"')
    h(d,'Задание для самостоятельной работы')
    d.paragraphs[-1].paragraph_format.page_break_before = True
    p(d,'Ситуация: линия 55 км, четыре разъёма по 0,4 дБ, три сварки по 0,1 дБ, затухание 0,22 дБ/км и запас 2,5 дБ. Сначала вручную посчитай L_физ и L_план. Затем проверь себя командой ниже. Объясни, почему плановый бюджет больше физического.')
    code(d,'python -c "from lab01.link_budget import Link; x=Link(distance_km=55,attenuation_db_km=0.22,connectors=4,connector_loss_db=0.4,splices=3,splice_loss_db=0.1,engineering_margin_db=2.5); print(x.physical_loss_db,x.planned_loss_db)"')
    h(d,'Разбор задания')
    p(d,'55 × 0,22 = 12,1 дБ; разъёмы 4 × 0,4 = 1,6 дБ; сварки 3 × 0,1 = 0,3 дБ. Физический бюджет 14,0 дБ, плановый — 16,5 дБ. Дополнительные 2,5 дБ нужны как резерв на старение, загрязнение и неопределённость проекта; их не следует выдавать за уже измеренную потерю.')
    h(d,'Практическая работа и собеседование')
    p(d,'В реальной задаче ты запросишь длину трассы, паспортное затухание волокна на рабочей длине волны, число соединений, измеренные потери и характеристики детектора. Сможешь подготовить таблицу бюджета и сказать, какое измерение следует сделать первым. На интервью покажи расчёт для 40 и 100 км, назови разницу между физической потерей и запасом и честно скажи, что оценка сырых кликов ещё не является скоростью секретного ключа.')
    sources(d,['ITU-T Y.3806: https://www.itu.int/rec/T-REC-Y.3806-202109-I — мониторинг скорости генерации и QBER.','Численные параметры в лаборатории иллюстративны. Для проекта с реальным оборудованием используют его паспорт и измерения линии.'])
    d.save(ROOT/'lab-01-optical-link'/'guide.docx')


def lab2():
    d=setup_doc('Протокол BB84 и обнаружение ошибок','Лабораторная 2 · просеивание базисов · шум и перехват')
    p(d,'Цель: собственными глазами увидеть, как из случайных битов и двух базисов получается просеянная строка, почему около половины позиций отбрасывают и почему перехват с повторной отправкой вызывает ошибки. Это учебная модель подготовки и измерения, а не генератор безопасных ключей.')
    h(d,'Шаг 1 Запусти три контрольных опыта')
    p(d,'Открой PowerShell в quantum-telecom-labs. Аргумент --pulses задаёт число отправок, --noise — вероятность переворота результата Боба, --intercept включает Еву, --seed делает случайную последовательность воспроизводимой. Без флагов программа запускает идеальный канал с 10 000 импульсами.')
    code(d,'python -m lab02.bb84\npython -m lab02.bb84 --pulses 100000 --noise 0.05\npython -m lab02.bb84 --pulses 100000 --intercept')
    picture(d,bb_terminal,'Снимок вывода с перехватом при фиксированном seed 2026')
    h(d,'Шаг 2 Проследи один импульс вручную')
    p(d,'Алиса выбирает бит 0 или 1 и базис 0 или 1. Боб независимо выбирает базис. Если базисы совпали, в идеальной модели он получает отправленный бит. Если нет, результат равновероятно 0 или 1. После передачи стороны открыто сравнивают только базисы, оставляя позиции совпадения. Поэтому из 100 000 отправок обычно остаётся около 50 000 битов.')
    table(d,['Случай','Результат в этой модели','Оставляем?'],[['Базисы совпали','Бит Алисы, если не было шума','Да'],['Базисы различны','Случайный бит Боба','Нет'],['Ева угадала базис','Повторно отправлен верный бит','Если базисы Алисы и Боба совпали'],['Ева ошиблась базисом','Измерение/повторная отправка дают риск ошибки','Если базисы Алисы и Боба совпали']])
    h(d,'Шаг 3 Пойми число около 25 процентов')
    p(d,'В простом сценарии intercept-resend Ева измеряет каждый импульс в случайном базисе и пересылает свой результат. На оставленных позициях она ошибается базисом примерно в половине случаев. В этих случаях Боб, измеряя в базисе Алисы, получает неверный бит примерно с вероятностью 1/2. Значит, ожидаемый QBER ≈ 1/2 × 1/2 = 25%. Конечная выборка может дать 24,96% или другую близкую величину.')
    code(d,'QBER = число_разных_битов_Алисы_и_Боба / число_просеянных_позиций')
    p(d,'Без перехвата и шума QBER должен быть 0. При --noise 0.05 он должен быть около 5%, но не строго равен 5%: программа моделирует случайные исходы. Параметр seed нужен для повторяемого эксперимента и тестов, а не для криптографической случайности.')
    h(d,'Шаг 4 Прочти ключевые строки программы')
    code(d,'a_bit, a_basis, b_basis = rng.randrange(2), rng.randrange(2), rng.randrange(2)\nif intercept:\n    eve_basis = rng.randrange(2)\n    transmitted_bit = measure(a_bit, a_basis, eve_basis, rng)\nif a_basis == b_basis:\n    alice.append(a_bit)\n    bob.append(b_bit)')
    p(d,'randrange(2) возвращает 0 или 1. Функция measure возвращает исходный бит при совпадении базисов и случайный при несовпадении. При включённом перехвате программа заменяет передаваемый импульс результатом измерения Евы. Условие a_basis == b_basis реализует просеивание: иначе пара битов не попадает в выборку для QBER.')
    h(d,'Шаг 5 Проведи собственный эксперимент')
    n(d,'Запусти --pulses 100, 10000 и 100000 с --intercept. Запиши число просеянных битов и QBER. Почему результат на 100 импульсах сильнее отклоняется от 25%?')
    n(d,'Запусти --noise 0.05 --intercept. Предскажи: QBER будет выше 25% или ниже? Проверь. Не складывай 25% и 5% напрямую — шум может как создать, так и исправить ошибку на конкретной позиции.')
    n(d,'Смени --seed 2026 на 7. Прогноз для большого n должен остаться примерно тем же, но конкретные биты и точный QBER изменятся.')
    code(d,'python -m lab02.bb84 --pulses 100 --intercept\npython -m lab02.bb84 --pulses 10000 --intercept\npython -m lab02.bb84 --pulses 100000 --intercept\npython -m lab02.bb84 --pulses 100000 --noise 0.05 --intercept\npython -m lab02.bb84 --pulses 100000 --intercept --seed 7')
    h(d,'Задание и разбор')
    p(d,'Задание: объясни человеку без квантовой подготовки, почему Алиса и Боб раскрывают базисы, но не публикуют всю просеянную строку. Затем ответь: если QBER 0%, доказана ли секретность?')
    p(d,'Разбор: публичное сравнение базисов позволяет выбрать измерения, сделанные в одной системе координат. Раскрытие самой строки сделало бы её известной наблюдателю. QBER 0% в коротком образце не доказывает секретность: модель не включает все атаки, статистическую оценку, аутентификацию, коррекцию ошибок и усиление секретности.')
    h(d,'Практическая работа и собеседование')
    p(d,'На интервью ты сможешь нарисовать цепочку «подготовка → измерение → раскрытие базисов → оценка QBER», объяснить происхождение 25% и показать эксперимент на 100 и 100 000 импульсах. Следующий профессиональный шаг — изучить конечные ключи, decoy-state BB84 и аппаратные несовершенства; эту работу нельзя выдавать за доказательство стойкости реального QKD.')
    sources(d,['ITU-T Y.3802: https://www.itu.int/rec/T-REC-Y.3802-202012-I — место QKD-модулей в сетевой архитектуре.','Исходный код лаборатории: ../lab02/bb84.py. В README перечислены отсутствующие этапы реального протокола.'])
    d.save(ROOT/'lab-02-bb84'/'guide.docx')


def lab3():
    d=setup_doc('Сервис выдачи ключей и диагностика отказов','Лабораторная 3 · REST API · жизненный цикл ключа · интеграция')
    p(d,'Цель: соединить представление о квантовом канале с обычной сетевой инженерией. Ты запустишь локальный менеджер ключей, запросишь один и тот же ключ для двух приложений, проверишь запрет повторного чтения и воспроизведёшь отказы. Источник ключей в стенде эмулирован; настоящий QKD-модуль к нему не подключён.')
    h(d,'Архитектура стенда')
    p(d,'Приложение Alice запрашивает материал у KeyStore через enc_keys и получает идентификатор key_ID. Приложение Bob запрашивает этот же материал через dec_keys по key_ID. KeyStore хранит пару приложений, ключ и доступный запас. HTTP-сервер принимает запросы только на 127.0.0.1. Это учебная версия идеи разделения приложения и менеджера ключей, описанной в ETSI GS QKD 014 и ITU-T Y.3802.')
    code(d,'Alice SAE ── enc_keys ──> локальный KME/KeyStore\nBob SAE   ── dec_keys ──> тот же KME/KeyStore\nQKD-оборудование отсутствует; начальный запас задаётся в памяти.')
    h(d,'Шаг 1 Запусти демонстрацию целиком')
    code(d,'python -m lab03.demo')
    p(d,'Команда -m запускает модуль demo. Он создаёт KeyStore с двумя ключами, запускает локальный HTTP-сервер на свободном порту, делает запросы от Alice и Bob, затем останавливает сервер. Настоящее значение UUID каждый раз меняется. Код не печатает сам секретный ключ.')
    picture(d,api_terminal,'Снимок реального запуска; случайный UUID заменён меткой UUID')
    h(d,'Шаг 2 Разбери маршруты и ответы')
    table(d,['Маршрут','Кто вызывает','Ожидаемый смысл'],[['GET /api/v1/keys/bob/status','Alice','Состояние источника, запас и размер ключа'],['GET /api/v1/keys/bob/enc_keys','Alice','Выдать Alice новый ключ для пары Alice–Bob'],['GET /api/v1/keys/alice/dec_keys?key_ID=...','Bob','Выдать Bob соответствующий ключ один раз']])
    p(d,'Заголовок X-SAE-ID указывает имя вызывающего приложения. В этом стенде он не удостоверяет личность: любой локальный клиент может подставить другое значение. Маршруты и поля key_ID/key демонстрируют структуру, похожую на ETSI GS QKD 014, но сервис не заявляет соответствие стандарту.')
    h(d,'Шаг 3 Запусти сервер отдельно и проверь запрос вручную')
    n(d,'В первом PowerShell выполни команду ниже. Оставь окно открытым; Ctrl+C остановит сервер.')
    code(d,'python -m lab03.service')
    n(d,'Во втором PowerShell запусти запрос состояния. Invoke-RestMethod преобразует JSON в объект PowerShell; -Headers передаёт учебный идентификатор Alice.')
    code(d,'Invoke-RestMethod -Uri "http://127.0.0.1:8080/api/v1/keys/bob/status" -Headers @{"X-SAE-ID"="alice"}')
    n(d,'Проверь, что source_online = True, qber = 0.02, stored_key_count = 4. Не публикуй результаты вызова enc_keys: ответ содержит материал ключа.')
    h(d,'Шаг 4 Прочти жизненный цикл одного ключа')
    p(d,'secrets.token_bytes создаёт случайные байты, base64 лишь кодирует их для передачи в JSON — это не шифрование. uuid4 создаёт идентификатор. При выдаче Alice запас уменьшается, а запись сохраняется с парой приложений. Bob получает её по ID; после выдачи запись удаляется. Повторный запрос возвращает HTTP 404. Блокировка threading.Lock защищает операции с общим словарём от одновременных запросов.')
    code(d,'key = base64.b64encode(secrets.token_bytes(32)).decode("ascii")\nself.issued[key_id] = (sender, receiver, key)\nself.capacity -= 1')
    h(d,'Шаг 5 Диагностируй четыре отказа')
    table(d,['Симптом','Причина в модели','Проверка'],[['HTTP 404 при повторном dec_keys','Ключ уже выдан и удалён','Посмотри issued после первого чтения'],['HTTP 503 при number > запас','Очередь исчерпана','Сравни number со stored_key_count'],['HTTP 503 при QBER > 0.11','Сработала локальная политика','Проверь qber и qber_limit'],['HTTP 503 при offline','Источник отключён','Проверь source_online']])
    p(d,'В demo.py qber и online намеренно меняются в памяти, чтобы воспроизвести последние два отказа. Ответ 503 означает, что сервис сейчас не может выдать ключ. В настоящей системе расследование дополнительно включало бы состояние оптического тракта, QKD-модуля, KME, сертификатов и связи между площадками.')
    h(d,'Задание для самостоятельной работы')
    p(d,'Открой lab03/demo.py и найди в main строку KeyStore(capacity=2). Измени её на KeyStore(capacity=1). До запуска предскажи, что изменится в initial status. Затем после строки print("repeat read rejected:", status), но до store.qber = 0.20, вставь две строки ниже. Повтори python -m lab03.demo: дополнительный запрос должен вернуть 503 из-за пустого запаса. После проверки верни capacity=2 и удали добавленные строки. Отсутствие HTTPS и настоящей аутентификации запрещает использовать стенд для защиты реальных данных.')
    code(d,'status, body = request(base, "bob", "enc_keys", "alice")\nprint("empty stock:", status, body["error"])')
    h(d,'Разбор задания')
    p(d,'Первый запрос Alice к enc_keys израсходует единственный доступный ключ; Bob по ID получит тот же материал из issued. Следующий запрос нового ключа получит HTTP 503 «insufficient key stock». HTTP передаёт данные без защиты канала, а X-SAE-ID можно подделать; поэтому демонстрация ограничена localhost и не является безопасным сервисом.')
    h(d,'Проверка и собеседование')
    code(d,'python -m unittest discover -s tests -v')
    p(d,'Команда discover находит тесты всех трёх лабораторий. Для этой работы она проверяет совпадение ключа у двух приложений, одноразовую выдачу, чужую пару, запас, QBER и отказ источника. На интервью покажи архитектуру, живой demo и последовательность «симптом → метрика → слой проблемы → исправление». Отдельно скажи, что для производства потребуются реальные QKD-устройства, взаимная аутентификация, TLS, защищённое хранение, управление жизненным циклом и аудит без записи секретов.')
    sources(d,['ETSI GS QKD 014: https://www.etsi.org/deliver/etsi_gs/QKD/001_099/014/01.01.01_60/gs_qkd014v010101p.pdf — API выдачи ключей.','ITU-T Y.3802: https://www.itu.int/rec/T-REC-Y.3802-202012-I — сетевая архитектура.','ITU-T Y.3806: https://www.itu.int/rec/T-REC-Y.3806-202109-I — метрики качества, включая QBER и задержку доставки.'])
    d.save(ROOT/'lab-03-key-service'/'guide.docx')


if __name__ == '__main__':
    lab1(); lab2(); lab3()
    print('Created three guides')
